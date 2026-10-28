"""附件按实际字节分流；不运行模型或访问微信。"""
import importlib.util
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from tool.content_reading import media


class FileDetectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'content.bin'
        self.output = Path(self.temp.name) / 'prepared'

    def prepare(self, name='misleading.txt'):
        return media.prepare_content(self.path, 'file', name=name, output_dir=self.output)

    def test_m4a_with_wrong_suffix_uses_audio_stream_not_cover_image(self):
        self.path.write_bytes(b'\x00\x00\x00\x20ftypM4A ' + b'\x00' * 40)
        info = b'  Stream #0:0: Audio: aac\n  Stream #0:1: Video: mjpeg (attached pic)\n'
        with patch.object(media, '_ffmpeg', return_value='ffmpeg'), \
                patch.object(media.subprocess, 'run', return_value=subprocess.CompletedProcess([], 1, b'', info)), \
                patch.object(media, '_audio', return_value={'transcript': 'test'}) as audio, \
                patch.object(media, '_video') as video:
            result = self.prepare()
        self.assertEqual(result['detected_type'], 'audio')
        self.assertEqual(result['transcript'], 'test')
        audio.assert_called_once_with(self.path.resolve(), self.output.resolve(), None)
        video.assert_not_called()

    def test_mp4_video_is_not_classified_as_audio(self):
        self.path.write_bytes(b'\x00\x00\x00\x20ftypisom' + b'\x00' * 40)
        info = b'  Stream #0:0: Audio: aac\n  Stream #0:1: Video: h264\n'
        with patch.object(media, '_ffmpeg', return_value='ffmpeg'), \
                patch.object(media.subprocess, 'run', return_value=subprocess.CompletedProcess([], 1, b'', info)), \
                patch.object(media, '_video', return_value={'images': []}) as video:
            self.assertEqual(self.prepare()['detected_type'], 'video')
        video.assert_called_once()

    @unittest.skipUnless(importlib.util.find_spec('imageio_ffmpeg'), 'optional ffmpeg dependency absent')
    def test_real_m4a_probe_and_conversion_without_model(self):
        source = Path(self.temp.name) / 'sample.m4a'
        media._run([media._ffmpeg(), '-nostdin', '-y', '-v', 'error', '-f', 'lavfi',
                    '-i', 'sine=frequency=440:duration=0.2', '-c:a', 'aac', str(source)])
        source.replace(self.path)
        with patch.object(media.transcription, 'transcribe_pcm', return_value={'transcript': '模拟转写'}) as recognize:
            result = self.prepare()
        self.assertEqual(result['detected_type'], 'audio')
        self.assertEqual(result['transcript'], '模拟转写')
        self.assertTrue(Path(result['audio']).is_file())
        recognize.assert_called_once()

    def test_pdf_signature_overrides_audio_name(self):
        self.path.write_bytes(b'%PDF-1.7\n')
        with patch.object(media, '_document', return_value={'text': 'pdf'}) as document:
            self.assertEqual(self.prepare('fake.m4a')['detected_type'], 'pdf')
        self.assertEqual(document.call_args.args[1], 'content.pdf')

    @unittest.skipUnless(importlib.util.find_spec('PIL'), 'optional image dependency absent')
    def test_image_without_extension_produces_visual_content(self):
        from PIL import Image
        Image.new('RGB', (20, 10), 'red').save(self.path, format='PNG')
        result = self.prepare('fake.pdf')
        self.assertEqual(result['detected_type'], 'image')
        self.assertTrue(Path(result['images'][0]).is_file())

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'optional spreadsheet dependency absent')
    def test_xlsx_without_extension_is_read_from_actual_container(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active['A1'] = '识别成功'
        book.save(self.path)
        book.close()
        result = self.prepare('wrong.txt')
        self.assertEqual(result['detected_type'], 'xlsx')
        self.assertIn('识别成功', result['text'])

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'optional document dependency absent')
    def test_docx_with_download_suffix_is_read_from_actual_container(self):
        from docx import Document
        document = Document()
        document.add_paragraph('正文证据')
        document.save(self.path)
        result = self.prepare('report.docx.1')
        self.assertEqual(result['detected_type'], 'docx')
        self.assertIn('正文证据', result['text'])

    def test_apk_and_corrupt_zip_have_specific_errors(self):
        with zipfile.ZipFile(self.path, 'w') as archive:
            archive.writestr('AndroidManifest.xml', b'example')
        with self.assertRaisesRegex(ValueError, 'APK'):
            self.prepare()
        self.path.write_bytes(b'PK\x03\x04broken')
        with self.assertRaisesRegex(ValueError, 'ZIP容器损坏'):
            self.prepare()

    def test_invalid_media_has_probe_reason_and_no_transcription(self):
        self.path.write_bytes(b'ID3invalid')
        with patch.object(media, '_ffmpeg', return_value='ffmpeg'), \
                patch.object(media.subprocess, 'run', return_value=subprocess.CompletedProcess([], 1, b'', b'bad')), \
                patch.object(media, '_audio') as audio:
            with self.assertRaisesRegex(ValueError, '未找到可读取的音频或视频流'):
                self.prepare()
        audio.assert_not_called()

    def test_text_named_like_pdf_is_not_sent_to_pdf_parser(self):
        self.path.write_text('普通文本', encoding='utf-8')
        self.assertEqual(self.prepare('fake.pdf')['text'], '普通文本')
