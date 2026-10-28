"""为任意代理提供有界的本地媒体和文档预处理。"""
import base64
import importlib.util
from io import BytesIO
from pathlib import Path
import re
import subprocess
import threading
import zipfile
from tool.settings import cache_directory
from tool.media_runtime import run_media as _run, ffmpeg_executable as _ffmpeg
from tool import transcription

MAX_TEXT = 100000
MAX_IMAGES = 8
ASR_MODEL = transcription.DEFAULT_MODEL_DIR
transcribe = transcription.transcribe_pcm
_pdf_lock = threading.Lock()


def dependencies(model_dir=None):
    model_dir = Path(model_dir) if model_dir else ASR_MODEL
    names = {'images': 'PIL', 'pdf': 'pypdf', 'pdf_render': 'pypdfium2', 'word': 'docx', 'spreadsheet': 'openpyxl',
             'audio_video': 'imageio_ffmpeg', 'local_transcription': 'faster_whisper'}
    return {**{key: importlib.util.find_spec(module) is not None for key, module in names.items()},
            'transcription_weights': (model_dir / 'model.bin').is_file()}


def _image(path, output):
    from PIL import Image
    with Image.open(path) as image:
        count = getattr(image, 'n_frames', 1)
        frames = sorted({round(i * (count - 1) / max(1, min(count, MAX_IMAGES) - 1))
                         for i in range(min(count, MAX_IMAGES))})
        files = []
        for index in frames:
            image.seek(index)
            frame = image.convert('RGB')
            frame.thumbnail((1600, 1600))
            target = output / f'frame-{index}.png'
            frame.save(target)
            files.append(str(target.resolve()))
        return {'images': files, 'dimensions': list(image.size), 'frame_count': count,
                'sampled_frames': frames}


def _audio(path, output, model_dir=None):
    return transcription.transcribe_audio(path, output_dir=output, model_dir=model_dir)


def detect_file(path):
    """以文件内容识别附件；扩展名只用于提示媒体探测，不作为类型证据。"""
    with path.open('rb') as stream:
        head = stream.read(4096)
    if not head:
        raise ValueError('附件为空')
    if head.startswith(b'%PDF-'):
        return 'document', '.pdf'
    if head.startswith((b'PK\x03\x04', b'PK\x05\x06', b'PK\x07\x08')):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            raise ValueError('ZIP容器损坏，无法识别附件类型') from None
        if '[Content_Types].xml' in names:
            if 'word/document.xml' in names:
                return 'document', '.docx'
            if 'xl/workbook.xml' in names:
                return 'document', '.xlsx'
        if 'AndroidManifest.xml' in names:
            raise ValueError('已识别为APK安装包，尚无APK内容解析器')
        raise ValueError('已识别为ZIP容器，尚无对应内容解析器')
    if head.startswith((b'\x89PNG\r\n\x1a\n', b'\xff\xd8\xff', b'GIF87a', b'GIF89a',
                        b'BM', b'II*\x00', b'MM\x00*', b'\x00\x00\x01\x00')) or (
            head.startswith(b'RIFF') and head[8:12] == b'WEBP'):
        return 'image', ''
    if head.startswith((b'MZ', b'\x7fELF', b'Rar!\x1a\x07', b'7z\xbc\xaf\x27\x1c',
                        b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')):
        raise ValueError('该二进制文件格式尚无内容解析器；不会执行附件')
    media_suffixes = {'.m4a', '.mp3', '.wav', '.flac', '.ogg', '.opus', '.aac', '.amr',
                      '.mp4', '.mov', '.mkv', '.webm', '.avi', '.mpeg', '.mpg', '.wma', '.wmv'}
    is_media = (head.startswith((b'RIFF', b'ID3', b'fLaC', b'OggS', b'#!AMR',
                                 b'\x1a\x45\xdf\xa3', b'\x30\x26\xb2\x75', b'\x00\x00\x01\xba')) or
                head[4:8] == b'ftyp' or (len(head) > 1 and head[0] == 255 and head[1] & 224 == 224))
    if is_media or path.suffix.lower() in media_suffixes:
        probe = subprocess.run([_ffmpeg(), '-nostdin', '-hide_banner', '-protocol_whitelist',
                                'file,pipe', '-i', str(path)], stdin=subprocess.DEVNULL,
                               capture_output=True, timeout=15,
                               creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        info = probe.stderr.decode('utf-8', errors='replace')
        # ffmpeg仅探测、不指定输出时退出码非零；以实际流信息判断。
        streams = [line for line in info.splitlines() if re.match(r'\s*Stream #', line)]
        if any('Video:' in line and '(attached pic)' not in line for line in streams):
            return 'video', ''
        if any('Audio:' in line for line in streams):
            return 'audio', ''
        raise ValueError('媒体类型探测失败：未找到可读取的音频或视频流')
    return 'document', ''


def _video(path, output, model_dir=None):
    # ffmpeg 同时负责元数据探测和画面解码，不依赖全局 PATH。
    probe = subprocess.run([_ffmpeg(), '-nostdin', '-i', str(path)], stdin=subprocess.DEVNULL,
                           capture_output=True, timeout=15, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    info = probe.stderr.decode('utf-8', errors='replace')
    match = re.search(r'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)', info)
    if not match:
        raise ValueError('无法验证视频时长')
    duration = int(match[1]) * 3600 + int(match[2]) * 60 + float(match[3])
    if not 0 < duration <= 600:
        raise ValueError('视频时长须在10分钟以内')
    _run([_ffmpeg(), '-nostdin', '-y', '-v', 'error', '-i', str(path), '-an',
          '-vf', f'fps={MAX_IMAGES / duration:.9f},scale=960:960:force_original_aspect_ratio=decrease',
          '-frames:v', str(MAX_IMAGES), '-q:v', '4', str(output / 'video-%02d.jpg')])
    images = [str(p.resolve()) for p in sorted(output.glob('video-*.jpg'))]
    if not images:
        raise ValueError('视频没有可解码画面')
    result = {'images': images, 'duration_seconds': duration,
              'sampled_frame_seconds_approx': [round((i + .5) * duration / MAX_IMAGES, 3) for i in range(len(images))],
              'visual_bytes': sum(Path(p).stat().st_size for p in images),
              'sampling': '全片均匀抽取至多8帧；不能据此推断未采样瞬间'}
    if 'Audio:' in info:
        try:
            result.update(_audio(path, output, model_dir))
        except ValueError as error:
            result['audio_error'] = str(error)
    return result


def _document(path, name, output=None):
    suffix = Path(name).suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(path)
        if pdf.is_encrypted:
            raise ValueError('PDF已加密，需要先提供可读取版本')
        pages = pdf.pages[:100]
        extracted = [page.extract_text() or '' for page in pages]
        text = '\n'.join(f'第{i + 1}页\n{value}' for i, value in enumerate(extracted))
        if not any(value.strip() for value in extracted):
            if output is None:
                raise ValueError('扫描PDF需要视觉处理输出目录')
            import pypdfium2 as pdfium
            # PDFium 不允许工作线程并发调用。
            with _pdf_lock:
                document = pdfium.PdfDocument(str(path))
                try:
                    count = len(document)
                    if count == 0:
                        raise ValueError('PDF没有页面')
                    selected = sorted({round(i * (count - 1) / max(1, min(count, MAX_IMAGES) - 1))
                                       for i in range(min(count, MAX_IMAGES))})
                    images = []
                    for index in selected:
                        page = document[index]
                        bitmap = None
                        try:
                            width, height = page.get_size()
                            if min(width, height) <= 0:
                                raise ValueError('PDF页面尺寸无效')
                            bitmap = page.render(scale=min(2, 1600 / max(width, height)), may_draw_forms=False)
                            image = bitmap.to_pil().convert('RGB')
                            target = output / f'pdf-page-{index + 1}.png'
                            image.save(target)
                            image.close()
                            images.append(str(target.resolve()))
                        finally:
                            if bitmap is not None:
                                bitmap.close()
                            page.close()
                finally:
                    document.close()
            return {'images': images, 'page_count': count, 'sampled_pages': [i + 1 for i in selected],
                    'truncated': len(selected) < count,
                    'sampling': '扫描PDF，至多8页；仅可根据已提供页面阅读，不能声称读完未提供的页面'}
        return {'text': text[:MAX_TEXT], 'truncated': len(pdf.pages) > 100 or len(text) > MAX_TEXT}
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        parts = [p.text for p in document.paragraphs]
        parts.extend('\t'.join(c.text for c in row.cells) for table in document.tables for row in table.rows)
        text = '\n'.join(parts)
        return {'text': text[:MAX_TEXT], 'truncated': len(text) > MAX_TEXT}
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=False)
        try:
            parts = []
            truncated = len(workbook.worksheets) > 20
            for sheet in workbook.worksheets[:20]:
                parts.append('工作表：' + sheet.title)
                truncated |= sheet.max_row > 1000 or sheet.max_column > 100
                for row in sheet.iter_rows(max_row=min(sheet.max_row, 1000), max_col=min(sheet.max_column, 100), values_only=True):
                    parts.append('\t'.join('' if v is None else str(v) for v in row))
            text = '\n'.join(parts)
            return {'text': text[:MAX_TEXT], 'truncated': truncated or len(text) > MAX_TEXT}
        finally:
            workbook.close()
    raw = path.read_bytes()
    if b'\0' in raw[:4096] or suffix in {'.exe', '.dll', '.zip', '.rar', '.7z', '.doc', '.xls'}:
        raise ValueError('该二进制文件格式尚无内容解析器；不会执行附件')
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            text = raw.decode(encoding)
            return {'text': text[:MAX_TEXT], 'encoding': encoding, 'truncated': len(text) > MAX_TEXT}
        except UnicodeDecodeError:
            continue
    raise ValueError('无法确定文本编码')


def prepare_content(path, kind, *, name=None, output_dir=None, model_dir=None):
    """读取本地文件，输出提取出的证据，不包含模型解读。"""
    try:
        path = Path(path).resolve(strict=True)
        if not path.is_file():
            raise ValueError('输入必须为本机文件')
        output = Path(output_dir).resolve() if output_dir else cache_directory('content-reading-')
        output.mkdir(parents=True, exist_ok=True)
        detected = None
        document_name = name or path.name
        if kind in ('document', 'file'):
            kind, suffix = detect_file(path)
            detected = suffix.lstrip('.') or kind
            # 使用内容识别出的文档格式，避免伪装/重复下载后缀影响解析。
            document_name = 'content' + suffix if suffix else 'content.txt'
        if kind in ('image', 'sticker'):
            result = _image(path, output)
        elif kind in ('audio', 'voice'):
            result = _audio(path, output, model_dir)
        elif kind == 'video':
            result = _video(path, output, model_dir)
        elif kind in ('document', 'file'):
            result = _document(path, document_name, output)
        else:
            raise ValueError('未知内容类型')
        if detected:
            result = {**result, 'detected_type': detected}
        return result
    except ValueError:
        raise
    except ImportError as error:
        raise ValueError('缺少内容处理依赖：' + str(error.name)) from None
    except Exception as error:
        reason = error.strerror if isinstance(error, OSError) else str(error)
        raise ValueError('内容解析失败：' + type(error).__name__ + '：' + str(reason or '未知原因')[:500]) from None


def external_content(text, attachment=None):
    """Responses 函数输出内容保留 ExternalMessage 的权限级别。"""
    if not attachment or not attachment.get('prepared'):
        return text
    content = [{'type': 'input_text', 'text': text}]
    prepared = attachment['prepared']
    for path in prepared.get('images', []):
        mime = 'image/jpeg' if Path(path).suffix.lower() in ('.jpg', '.jpeg') else 'image/png'
        content.append({'type': 'input_image', 'image_url': 'data:' + mime + ';base64,' +
                        base64.b64encode(Path(path).read_bytes()).decode()})
    if prepared.get('audio') and 'transcript' not in prepared:
        content.append({'type': 'input_audio', 'audio_url': 'data:audio/wav;base64,' +
                        base64.b64encode(Path(prepared['audio']).read_bytes()).decode()})
    return content
